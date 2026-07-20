#!/usr/bin/env python3
"""
Build the KC Resources dataset.

resources.csv  ->  resources.xlsx   (human-editable master)
               ->  resources.json   (bundled into the iOS app + seeded into the server)

Usage:
    python3 build_dataset.py                # convert only
    python3 build_dataset.py --geocode      # also fill in missing lat/lng via OpenStreetMap

--geocode requires internet and writes coordinates back into resources.csv so you
only pay the geocoding cost once. Nominatim asks for max 1 request/second; we honor that.

Entries whose address is a placeholder ("Phone only", "Confidential - call hotline",
"Mobile - varies by night", "Call for current address") are intentionally left without
coordinates. The app renders those as call-only cards with no map pin.
"""

import argparse
import csv
import json
import re
import sys
import time
from pathlib import Path

HERE = Path(__file__).parent
CSV_PATH = HERE / "resources.csv"
JSON_PATH = HERE / "resources.json"
XLSX_PATH = HERE / "resources.xlsx"
IOS_JSON = HERE.parent / "ios/KCResources/KCResources/Resources/resources.json"

# Addresses that are deliberately not mappable.
NON_ADDRESS = re.compile(
    r"^(phone only|confidential|mobile\s*-|call for current address)", re.IGNORECASE
)

# Loose bounding box for the KC metro. Anything outside is almost certainly a bad geocode.
KC_BOUNDS = {"lat": (38.6, 39.5), "lon": (-95.1, -94.2)}


def is_mappable(address: str) -> bool:
    return bool(address) and not NON_ADDRESS.match(address.strip())


# Unit qualifiers confuse Nominatim — "14 W 10th Street, Third Floor" returns
# nothing while "14 W 10th Street" resolves fine. Same for suite and apartment
# numbers. We only need building-level accuracy anyway.
UNIT_SUFFIX = re.compile(
    r",?\s*(?:"
    r"#\s*\S+"                     # "#14" — \b won't match before '#'
    r"|\b(?:suite|ste\.?|apt\.?|apartment|unit|floor|fl\.?|"
    r"(?:first|second|third|fourth|fifth|ground)\s+floor)\b.*"
    r")$",
    re.IGNORECASE,
)


def geocodable(address: str) -> str:
    """
    Normalises a street address into something Nominatim can resolve.

    Two known failure modes in this dataset:
      "14 W 10th Street, Third Floor"  -> drop the floor
      "40th & Main Street"             -> intersections want 'and', not '&'
    """
    cleaned = UNIT_SUFFIX.sub("", address.strip()).strip(" ,")
    if "&" in cleaned:
        cleaned = cleaned.replace("&", "and")
    return cleaned


def read_rows():
    with CSV_PATH.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def write_rows(rows):
    if not rows:
        return
    with CSV_PATH.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def geocode(rows):
    import urllib.parse
    import urllib.request

    filled = skipped = failed = 0
    for row in rows:
        if row.get("latitude") and row.get("longitude"):
            continue
        if not is_mappable(row["address"]):
            skipped += 1
            continue

        query = ", ".join(
            p for p in [geocodable(row["address"]), row["city"], row["state"], row["zip_code"]] if p
        )
        url = "https://nominatim.openstreetmap.org/search?" + urllib.parse.urlencode(
            {"format": "json", "limit": 1, "countrycodes": "us", "q": query}
        )
        req = urllib.request.Request(url, headers={"User-Agent": "KCResourceTracker/1.0"})

        try:
            with urllib.request.urlopen(req, timeout=20) as resp:
                data = json.load(resp)
        except Exception as e:  # noqa: BLE001
            print(f"  ! {row['name']}: request failed ({e})")
            failed += 1
            time.sleep(1.1)
            continue

        if not data:
            print(f"  ! {row['name']}: no match for '{query}'")
            failed += 1
        else:
            lat, lon = float(data[0]["lat"]), float(data[0]["lon"])
            lo, hi = KC_BOUNDS["lat"]
            lo2, hi2 = KC_BOUNDS["lon"]
            if not (lo <= lat <= hi and lo2 <= lon <= hi2):
                print(f"  ! {row['name']}: {lat},{lon} falls outside the KC metro - skipping")
                failed += 1
            else:
                row["latitude"], row["longitude"] = f"{lat:.6f}", f"{lon:.6f}"
                print(f"  + {row['name']}: {lat:.5f}, {lon:.5f}")
                filled += 1

        time.sleep(1.1)  # Nominatim usage policy

    print(f"\nGeocoded {filled}, skipped {skipped} non-address entries, {failed} failed.")
    return rows


def to_records(rows):
    records = []
    for row in rows:
        rec = {
            "id": row["id"],
            "name": row["name"],
            "organization": row["organization"],
            "category": row["category"],
            "services": [s.strip() for s in row["services"].split(";") if s.strip()],
            "description": row["description"],
            "address": row["address"],
            "city": row["city"],
            "state": row["state"],
            "zipCode": row["zip_code"],
            "phone": row["phone"],
            "website": row["website"],
            "hours": row["hours"],
            "eligibility": row["eligibility"],
            "verified": row["verified"].strip().upper() == "TRUE",
            "source": row["source"],
            "mappable": is_mappable(row["address"]),
        }
        if row.get("latitude") and row.get("longitude"):
            rec["latitude"] = float(row["latitude"])
            rec["longitude"] = float(row["longitude"])
        records.append(rec)
    return records


def write_xlsx(rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed - skipping .xlsx (pip install openpyxl)")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Resources"
    headers = list(rows[0].keys())
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E5F")
    for i, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(vertical="center")

    unverified_fill = PatternFill("solid", fgColor="FFF3CD")
    for row in rows:
        ws.append([row[h] for h in headers])
        if row["verified"].strip().upper() != "TRUE":
            for i in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=i).fill = unverified_fill

    widths = {
        "name": 42, "organization": 30, "description": 60, "address": 32,
        "hours": 46, "eligibility": 30, "services": 34, "website": 40, "source": 28,
    }
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    notes = wb.create_sheet("How to use")
    for line in [
        ["Editing this dataset"],
        [""],
        ["1. Edit rows on the Resources sheet, or edit resources.csv directly."],
        ["2. Rows shaded yellow have verified = FALSE. Confirm them by phone before trusting them."],
        ["3. Leave latitude/longitude blank for new rows."],
        ["4. Re-export:  python3 build_dataset.py --geocode"],
        ["   That fills in coordinates and regenerates resources.json for the app."],
        [""],
        ["Category must be one of: food, shelter, health, crisis"],
        ["Services is a semicolon-separated list, e.g.  Hot Meals;Showers;Mail"],
        ["Set address to 'Phone only' or 'Confidential - call hotline' for entries with no"],
        ["public location. Those show in the list as call-only and get no map pin."],
    ]:
        notes.append(line)
    notes["A1"].font = Font(bold=True, size=14)
    notes.column_dimensions["A"].width = 96

    wb.save(XLSX_PATH)
    print(f"Wrote {XLSX_PATH.name}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--geocode", action="store_true", help="fill missing coordinates via OpenStreetMap")
    args = ap.parse_args()

    if not CSV_PATH.exists():
        sys.exit(f"Missing {CSV_PATH}")

    rows = read_rows()

    if args.geocode:
        print("Geocoding (about 1 second per address)...\n")
        rows = geocode(rows)
        write_rows(rows)
        print(f"Updated {CSV_PATH.name} with coordinates.")

    records = to_records(rows)
    payload = {
        "version": 1,
        "generated": time.strftime("%Y-%m-%d"),
        "region": "Kansas City metro (MO + KS)",
        "resources": records,
    }
    JSON_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"Wrote {JSON_PATH.name} ({len(records)} resources)")

    if IOS_JSON.parent.exists():
        IOS_JSON.write_text(json.dumps(payload, indent=2), encoding="utf-8")
        print(f"Copied into the iOS bundle at {IOS_JSON.relative_to(HERE.parent)}")

    write_xlsx(rows)

    missing = [r["name"] for r in records if r["mappable"] and "latitude" not in r]
    if missing:
        print(f"\n{len(missing)} mappable entries still have no coordinates.")
        print("Run with --geocode, or let the app geocode them on first launch.")


if __name__ == "__main__":
    main()
